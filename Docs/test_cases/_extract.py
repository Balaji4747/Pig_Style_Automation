import openpyxl, csv, os
from collections import Counter

wb = openpyxl.load_workbook(r'docs/Pig_Style_QA_Test_Suite (1).xlsx', data_only=True)

TARGET = {
    'Interruption & Recovery': 'ps_int',
    'Error Handling': 'ps_err',
    'STR - Transaction Testing': 'ps_str',
    'Math Validation': 'ps_mth',
    'Feature Testing': 'ps_fea',
    'Edge & Boundary Cases': 'ps_edg',
    'UI-UX Testing': 'ps_ui',
    'Compatibility Testing': 'ps_cmp',
    'Regression Suite': 'ps_reg',
}
REFERENCE = {
    'Functional Testing': 'functional_reference',
    'Defect-Observation Log': 'defect_observation_log',
}
EXTRA_SHEETS = ['Test Case Summary', 'Traceability Matrix', 'Test Execution Summary', 'Math Reference']

outdir = 'docs/test_cases'
index_lines = [
    '# Test Case Workbook Index\n\n',
    'Source: `docs/Pig_Style_QA_Test_Suite (1).xlsx`\n\n',
    '| Sheet | Rows (incl. header) | Data Rows | Files |\n|---|---|---|---|\n',
]


def dump_sheet(sheet_name, slug):
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    data_rows = rows[1:]
    data_rows = [r for r in data_rows if any(v not in (None, '') for v in r)]
    csv_path = '{}/{}.csv'.format(outdir, slug)
    with open(csv_path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in data_rows:
            w.writerow(['' if v is None else v for v in r])
    return header, data_rows, csv_path


ALL_SHEETS = {}
ALL_SHEETS.update(TARGET)
ALL_SHEETS.update(REFERENCE)

for sheet_name, slug in ALL_SHEETS.items():
    header, data_rows, csv_path = dump_sheet(sheet_name, slug)
    note = '' if sheet_name in TARGET else ' (reference only, not in requested scope)'
    index_lines.append('| {}{} | {} | {} | `{}` |\n'.format(sheet_name, note, len(data_rows) + 1, len(data_rows), csv_path))

for sheet_name in EXTRA_SHEETS:
    slug = sheet_name.lower().replace(' ', '_').replace('&', 'and').replace('-', '_')
    header, data_rows, csv_path = dump_sheet(sheet_name, slug)
    index_lines.append('| {} (reference only, not in requested scope) | {} | {} | `{}` |\n'.format(sheet_name, len(data_rows) + 1, len(data_rows), csv_path))

with open('{}/INDEX.md'.format(outdir), 'w', encoding='utf-8') as f:
    f.writelines(index_lines)

ID_COL = 'Test Case ID'
PRIORITY_COL = 'Priority'
AUTOCAND_COL = 'Automation Candidate'
DESC_COL = 'Test Scenario'

for sheet_name, slug in TARGET.items():
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    data_rows = [r for r in rows[1:] if any(v not in (None, '') for v in r)]
    idx = {h: i for i, h in enumerate(header)}

    id_i = idx.get(ID_COL)
    pri_i = idx.get(PRIORITY_COL)
    auto_i = idx.get(AUTOCAND_COL)
    desc_i = idx.get(DESC_COL)

    cross = Counter()
    for r in data_rows:
        pri = (r[pri_i] if pri_i is not None else None) or '(blank)'
        auto = (r[auto_i] if auto_i is not None else None) or '(blank)'
        cross[(pri, auto)] += 1

    lines = []
    lines.append('# {} - Summary\n\n'.format(sheet_name))
    lines.append('Columns exactly as found in workbook: {}\n\n'.format(header))
    lines.append('Total data rows: {}\n\n'.format(len(data_rows)))
    lines.append('## Priority x Automation Candidate breakdown\n\n')
    lines.append('| Priority | Automation Candidate | Count |\n|---|---|---|\n')
    for (pri, auto), cnt in sorted(cross.items(), key=lambda kv: str(kv[0])):
        lines.append('| {} | {} | {} |\n'.format(pri, auto, cnt))
    lines.append('\n## Full row list (ID | Priority | Automation Candidate | Scenario)\n\n')
    lines.append('| Test Case ID | Priority | Automation Candidate | Test Scenario |\n|---|---|---|---|\n')
    for r in data_rows:
        tid = r[id_i] if id_i is not None else ''
        pri = r[pri_i] if pri_i is not None else ''
        auto = r[auto_i] if auto_i is not None else ''
        desc = (r[desc_i] if desc_i is not None else '') or ''
        desc = str(desc).replace('\n', ' ')[:140]
        lines.append('| {} | {} | {} | {} |\n'.format(tid, pri, auto, desc))

    with open('{}/{}_summary.md'.format(outdir, slug), 'w', encoding='utf-8') as f:
        f.writelines(lines)

print('DONE')
